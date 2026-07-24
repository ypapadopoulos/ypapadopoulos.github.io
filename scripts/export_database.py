#!/usr/bin/env python3

import argparse
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data-source"
OUTPUT_DIR = ROOT / "data"

DATABASES = {
    "theses": {
        "input": SOURCE_DIR / "theses.xlsx",
        "output": OUTPUT_DIR / "theses.yaml",
        "sheet": "Theses",
        "table": "ThesesTable",
    },
    "courses": {
        "input": SOURCE_DIR / "courses.xlsx",
        "output": OUTPUT_DIR / "courses.yaml",
        "sheet": "Courses",
        "table": "CoursesTable",
    },
}

ALLOWED_THESIS_TYPES = {"diploma", "masters", "phd"}
ALLOWED_STATUSES = {"draft", "submitted", "published", "archived"}
EXPORTED_STATUSES = {"submitted", "published"}

THESIS_OPTIONAL_FIELDS = [
    "pdf",
    "github_repository",
    "presentation_url",
    "publication_url",
    "repository_url",
    "featured_image",
]

COURSE_OPTIONAL_FIELDS = [
    "external_link",
    "image_filename",
]

COURSE_TRANSLATED_FIELDS = [
    "title",
    "summary",
    "semester",
    "level",
    "location",
    "schedule",
    "description",
    "teaching_schedule",
    "image_alt",
    "image_caption",
]


def text(value):
    if value is None:
        return ""

    return (
        str(value)
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )


def iso_date(value, row_number):
    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    value = text(value)

    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass

    raise ValueError(f"Row {row_number}: invalid date {value!r}")


def read_excel_table(config):
    workbook = load_workbook(config["input"], data_only=True)
    worksheet = workbook[config["sheet"]]
    table = worksheet.tables[config["table"]]
    cells = worksheet[table.ref]

    headers = [text(cell.value) for cell in cells[0]]
    rows = []

    for excel_row in cells[1:]:
        values = [cell.value for cell in excel_row]

        if all(text(value) == "" for value in values):
            continue

        rows.append((excel_row[0].row, dict(zip(headers, values))))

    return rows

def string_representer(dumper, value):
    style = "|" if "\n" in value else None
    return dumper.represent_scalar(
        "tag:yaml.org,2002:str",
        value,
        style=style,
    )


yaml.SafeDumper.add_representer(str, string_representer)


def write_yaml(records, output):
    output.parent.mkdir(parents=True, exist_ok=True)

    with output.open("w", encoding="utf-8") as file:
        for i, record in enumerate(records):
            yaml.safe_dump(
                [record],
                file,
                allow_unicode=True,
                sort_keys=False,
                width=1000,
            )

            if i < len(records) - 1:
                file.write("\n")


def export_theses(config):
    rows = read_excel_table(config)

    records = []
    seen_ids = set()
    seen_slugs = set()

    for row_number, row in rows:
        thesis_id = text(row.get("id"))
        slug = text(row.get("slug"))
        thesis_type = text(row.get("thesis_type")).lower()
        status = text(row.get("status")).lower()

        if not thesis_id:
            raise ValueError(f"Row {row_number}: missing id")

        if thesis_id in seen_ids:
            raise ValueError(
                f"Row {row_number}: duplicate id {thesis_id!r}"
            )

        if not slug:
            raise ValueError(f"Row {row_number}: missing slug")

        if slug in seen_slugs:
            raise ValueError(
                f"Row {row_number}: duplicate slug {slug!r}"
            )

        if thesis_type not in ALLOWED_THESIS_TYPES:
            raise ValueError(
                f"Row {row_number}: invalid thesis_type {thesis_type!r}"
            )

        if status not in ALLOWED_STATUSES:
            raise ValueError(
                f"Row {row_number}: invalid status {status!r}"
            )

        seen_ids.add(thesis_id)
        seen_slugs.add(slug)

        if status not in EXPORTED_STATUSES:
            continue

        record = {
            "id": thesis_id,
            "slug": slug,
            "date": iso_date(row.get("date"), row_number),
            "thesis_type": thesis_type,
            "status": status,
            "student": {
                "en": text(row.get("student_en")),
                "el": text(row.get("student_el")),
            },
            "title": {
                "en": text(row.get("title_en")),
                "el": text(row.get("title_el")),
            },
            "programme": {
                "en": text(row.get("programme_en")),
                "el": text(row.get("programme_el")),
            },
            "summary": {
                "en": text(row.get("summary_en")),
                "el": text(row.get("summary_el")),
            },
        }

        for field in THESIS_OPTIONAL_FIELDS:
            value = text(row.get(field))
            if value:
                record[field] = value

        records.append(record)

    records.sort(key=lambda item: item["date"], reverse=True)
    write_yaml(records, config["output"])

    print(
        f"Exported {len(records)} theses "
        f"to {config['output'].relative_to(ROOT)}"
    )


def export_courses(config):
    rows = read_excel_table(config)

    records = []
    seen_ids = set()
    seen_slugs = set()

    for row_number, row in rows:
        course_id = text(row.get("id"))
        slug = text(row.get("slug"))
        status = text(row.get("status")).lower()
        course_code = text(row.get("course_code"))
        ects = text(row.get("ects"))

        if not course_id:
            raise ValueError(f"Row {row_number}: missing id")

        if course_id in seen_ids:
            raise ValueError(
                f"Row {row_number}: duplicate id {course_id!r}"
            )

        if not slug:
            raise ValueError(f"Row {row_number}: missing slug")

        if slug in seen_slugs:
            raise ValueError(
                f"Row {row_number}: duplicate slug {slug!r}"
            )

        if status not in ALLOWED_STATUSES:
            raise ValueError(
                f"Row {row_number}: invalid status {status!r}"
            )

        seen_ids.add(course_id)
        seen_slugs.add(slug)

        if status not in EXPORTED_STATUSES:
            continue

        record = {
            "id": course_id,
            "slug": slug,
            "date": iso_date(row.get("date"), row_number),
            "status": status,
            "course_code": course_code,
            "ects": ects,
        }

        for field in COURSE_TRANSLATED_FIELDS:
            en = text(row.get(f"{field}_en"))
            el = text(row.get(f"{field}_el"))

            if en or el:
                record[field] = {
                    "en": en,
                    "el": el,
                }

        for field in COURSE_OPTIONAL_FIELDS:
            value = text(row.get(field))
            if value:
                record[field] = value

        tags = text(row.get("tags"))
        if tags:
            record["tags"] = [
                tag.strip()
                for tag in tags.split(";")
                if tag.strip()
            ]

        records.append(record)

    records.sort(key=lambda item: item["date"], reverse=True)
    write_yaml(records, config["output"])

    print(
        f"Exported {len(records)} courses "
        f"to {config['output'].relative_to(ROOT)}"
    )


EXPORTERS = {
    "theses": export_theses,
    "courses": export_courses,
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "database",
        choices=[*DATABASES, "all"],
    )
    args = parser.parse_args()

    names = DATABASES if args.database == "all" else [args.database]

    for name in names:
        EXPORTERS[name](DATABASES[name])


if __name__ == "__main__":
    main()